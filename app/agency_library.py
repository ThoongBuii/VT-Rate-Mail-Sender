"""Thư viện agency theo Country — lưu trong app, chọn country thay vì import mỗi lần."""

from __future__ import annotations

import json
import re
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .importer import _map_headers, _rows_to_mails
from .models import AgencyMail, MailStatus
from .paths import user_data_dir

LIBRARY_NAME = "agency_library.json"
SKIP_SHEETS = {"huong_dan", "hướng dẫn", "guide", "readme", "notes"}


def library_path() -> Path:
    return user_data_dir() / LIBRARY_NAME


def _norm_country(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "").strip())


def empty_library() -> dict[str, Any]:
    return {
        "source_name": "",
        "updated_at": 0.0,
        "active_country": "",
        "countries": {},  # name -> list[dict]
    }


def load_library() -> dict[str, Any]:
    path = library_path()
    if not path.exists():
        return empty_library()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return empty_library()
    if not isinstance(data, dict):
        return empty_library()
    countries = data.get("countries") or {}
    if not isinstance(countries, dict):
        countries = {}
    return {
        "source_name": str(data.get("source_name") or ""),
        "updated_at": float(data.get("updated_at") or 0),
        "active_country": str(data.get("active_country") or ""),
        "countries": countries,
    }


def save_library(data: dict[str, Any]) -> None:
    path = library_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source_name": str(data.get("source_name") or ""),
        "updated_at": float(data.get("updated_at") or time.time()),
        "active_country": str(data.get("active_country") or ""),
        "countries": data.get("countries") or {},
    }
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def mail_to_lib_dict(mail: AgencyMail) -> dict[str, Any]:
    return {
        "agency_company": mail.agency_company,
        "account_name": mail.account_name,
        "account_mail": mail.account_mail,
        "mail_cc": mail.mail_cc,
        "country": getattr(mail, "country", "") or "",
    }


def lib_dict_to_mail(item: dict[str, Any], row_index: int = 0) -> AgencyMail:
    return AgencyMail(
        agency_company=str(item.get("agency_company") or ""),
        account_name=str(item.get("account_name") or ""),
        account_mail=str(item.get("account_mail") or ""),
        mail_cc=str(item.get("mail_cc") or ""),
        country=str(item.get("country") or ""),
        status=MailStatus.PENDING,
        row_index=row_index,
    )


def country_names(lib: dict[str, Any]) -> list[str]:
    names = sorted((_norm_country(k) for k in (lib.get("countries") or {}) if _norm_country(k)))
    return names


def get_country_mails(lib: dict[str, Any], country: str) -> list[AgencyMail]:
    key = _norm_country(country)
    raw = (lib.get("countries") or {}).get(key) or []
    mails: list[AgencyMail] = []
    for i, item in enumerate(raw, start=1):
        if not isinstance(item, dict):
            continue
        mail = lib_dict_to_mail(item, row_index=i)
        mail.country = key
        mails.append(mail)
    return mails


def set_country_mails(lib: dict[str, Any], country: str, mails: list[AgencyMail]) -> dict[str, Any]:
    key = _norm_country(country)
    if not key:
        return lib
    countries = dict(lib.get("countries") or {})
    countries[key] = [mail_to_lib_dict(m) for m in mails]
    lib = dict(lib)
    lib["countries"] = countries
    lib["active_country"] = key
    lib["updated_at"] = time.time()
    return lib


def _sheet_to_mails(headers: list[str], rows: list[list], country: str) -> list[AgencyMail]:
    mails = _rows_to_mails(headers, rows)
    for m in mails:
        m.country = country
    return mails


def import_excel_to_library(path: str | Path, *, source_name: str = "") -> dict[str, Any]:
    """
    Nạp Excel vào thư viện:
    - Nhiều sheet → mỗi sheet = 1 country (trừ sheet hướng dẫn)
    - 1 sheet có cột Country → nhóm theo cột
    - 1 sheet không có Country → country = tên sheet / 'General'
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    by_country: dict[str, list[dict[str, Any]]] = {}

    for ws in wb.worksheets:
        title = _norm_country(ws.title or "")
        if title.lower() in SKIP_SHEETS:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(c or "") for c in rows[0]]
        mapping = _map_headers(headers)
        data_rows = [list(r) for r in rows[1:]]

        # Có cột Country → nhóm theo giá trị
        country_idx = mapping.get("country")
        if country_idx is not None:
            buckets: dict[str, list[list]] = {}
            for values in data_rows:
                if all(v is None or str(v).strip() == "" for v in values):
                    continue
                cname = _norm_country(str(values[country_idx] if country_idx < len(values) else "") or title or "General")
                if not cname:
                    cname = title or "General"
                buckets.setdefault(cname, []).append(values)
            for cname, brow in buckets.items():
                mails = _sheet_to_mails(headers, brow, cname)
                by_country.setdefault(cname, [])
                by_country[cname].extend(mail_to_lib_dict(m) for m in mails)
            continue

        # Không có cột Country → dùng tên sheet
        cname = title or "General"
        try:
            mails = _sheet_to_mails(headers, data_rows, cname)
        except ValueError:
            continue
        by_country.setdefault(cname, [])
        by_country[cname].extend(mail_to_lib_dict(m) for m in mails)

    if not by_country:
        raise ValueError(
            "Không đọc được agency nào.\n"
            "Excel: mỗi sheet = 1 country, hoặc có cột Country.\n"
            "Cột bắt buộc: Account Mail (+ Agency Company, Account Name, Mail cc)."
        )

    names = sorted(by_country.keys())
    active = names[0]
    lib = {
        "source_name": source_name or path.name,
        "updated_at": time.time(),
        "active_country": active,
        "countries": by_country,
    }
    save_library(lib)
    return lib


def library_summary(lib: dict[str, Any] | None = None) -> dict[str, Any]:
    lib = lib if lib is not None else load_library()
    countries = country_names(lib)
    counts = {
        name: len((lib.get("countries") or {}).get(name) or [])
        for name in countries
    }
    return {
        "source_name": lib.get("source_name") or "",
        "updated_at": lib.get("updated_at") or 0,
        "active_country": lib.get("active_country") or "",
        "countries": countries,
        "counts": counts,
        "has_library": bool(countries),
    }
